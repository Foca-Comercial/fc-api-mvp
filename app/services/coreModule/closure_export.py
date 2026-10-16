"""Genera el Excel (.xlsx) del cierre del día para cargar en Siigo.

Formato provisional (Siigo aún no define plantilla): un renglón por
vendedor + producto + precio, un total por vendedor y el total del día.
No depende de la BD: recibe las facturas ya consolidadas por el servicio.
"""

from collections.abc import Sequence
from datetime import date
from decimal import Decimal
from io import BytesIO
from typing import Protocol

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEADER = (
    "Fecha",
    "Cód. vendedor",
    "Vendedor",
    "Cód. producto",
    "Producto",
    "Cantidad",
    "Precio unit.",
    "Subtotal",
)
WIDTHS = (12, 14, 28, 14, 30, 11, 14, 16)
DATE_FORMAT = "DD/MM/YYYY"
QTY_FORMAT = "#,##0"
MONEY_FORMAT = '"$"#,##0.00'
MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

_BOLD = Font(bold=True)
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_TOTAL_FILL = PatternFill("solid", fgColor="DDEBF7")
_GRAND_FILL = PatternFill("solid", fgColor="BDD7EE")
_TOP_BORDER = Border(top=Side(style="thin"))


class _Line(Protocol):
    product_code: str
    product_name: str
    unit_price: Decimal
    quantity: int
    subtotal: Decimal


class _Seller(Protocol):
    seller_code: int
    seller_name: str

    @property
    def lines(self) -> Sequence[_Line]: ...

    @property
    def total_quantity(self) -> int: ...

    @property
    def total_sales(self) -> Decimal: ...


def filename_for(day: date) -> str:
    return f"cierre_{day.isoformat()}.xlsx"


def build_closure_workbook(day: date, sellers: Sequence[_Seller]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = f"Cierre {day.isoformat()}"  # type: ignore # máx. 31 caracteres

    ws.append(HEADER)  # type: ignore
    for cell in ws[1]:  # type: ignore
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for seller in sellers:
        for line in seller.lines:
            ws.append(  # type: ignore
                (
                    day,
                    seller.seller_code,  # número: Siigo recibe 1, no "01"
                    seller.seller_name,
                    line.product_code,
                    line.product_name,
                    line.quantity,
                    line.unit_price,
                    line.subtotal,
                )
            )
            _format_numbers(ws, ws.max_row)  # type: ignore
        ws.append(  # type: ignore
            (
                None,
                None,
                seller.seller_name,
                None,
                "Total vendedor",
                seller.total_quantity,
                None,
                seller.total_sales,
            )
        )
        _style_total(ws, ws.max_row, _TOTAL_FILL)  # type: ignore

    ws.append(  # type: ignore
        (
            None,
            None,
            None,
            None,
            "TOTAL DEL DÍA",
            sum(s.total_quantity for s in sellers),
            None,
            sum((s.total_sales for s in sellers), Decimal(0)),
        )
    )
    _style_total(ws, ws.max_row, _GRAND_FILL)  # type: ignore

    for index, width in enumerate(WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width  # type: ignore
    ws.freeze_panes = "A2"  # type: ignore

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _format_numbers(ws, row: int) -> None:
    ws.cell(row, 1).number_format = DATE_FORMAT
    ws.cell(row, 6).number_format = QTY_FORMAT
    ws.cell(row, 7).number_format = MONEY_FORMAT
    ws.cell(row, 8).number_format = MONEY_FORMAT


def _style_total(ws, row: int, fill: PatternFill) -> None:
    _format_numbers(ws, row)
    for cell in ws[row]:
        cell.font = _BOLD
        cell.fill = fill
        cell.border = _TOP_BORDER
