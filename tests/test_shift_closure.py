"""Vista de Cierre: CRUD de solo lectura (C1), resumen (C2) y exportación (C3)."""

from datetime import UTC, date, datetime, time, timedelta
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy.exc import IntegrityError

from app.core.config import settings
from app.core.timezone import BOGOTA_TZ, today_bogota
from app.models.coreModule.order import Order
from app.models.coreModule.shift_closure import ShiftClosure
from app.repositories.coreModule.shift_closure_repository import (
    shift_closure_repository,
)
from app.services.coreModule.closure_export import MEDIA_TYPE
from app.utils.enums.order_state import OrderState
from tests.helpers import API, item

URL = f"{API}/shifts_closures"


def test_el_crud_es_de_solo_lectura(admin):
    assert admin.get(f"{URL}/").status_code == 200
    assert admin.post(f"{URL}/", json={}).status_code == 405
    assert admin.patch(f"{URL}/1", json={}).status_code == 405
    assert admin.delete(f"{URL}/1").status_code == 405


def test_un_cierre_por_vendedor_y_dia(db, data):
    def closure():
        return ShiftClosure(
            closure_date=date(2026, 10, 15),
            id_user_vehicle=data.user_vehicle.id_user_vehicle,
            total_orders=1,
            total_sales=Decimal(1000),
        )

    db.add(closure())
    db.flush()
    db.add(closure())

    with pytest.raises(IntegrityError):
        db.flush()
    db.rollback()


# --- Resumen del cierre (C2): GET /shifts_closures/summary ---------------------------

SUMMARY = f"{URL}/summary"


def _dispatch(client, order: dict, deliver: bool = False) -> None:
    response = client.post(f"{API}/orders/{order['id_order']}/dispatch")
    assert response.status_code == 200, response.text
    if deliver:
        response = client.post(f"{API}/orders/{order['id_order']}/deliver")
        assert response.status_code == 200, response.text


def _move(db, order: dict, day: date) -> None:
    db.get(Order, order["id_order"]).creation_datetime = datetime.combine(
        day, time(10, 0), tzinfo=BOGOTA_TZ
    )
    db.commit()


def _summary(client, day: date | None = None) -> dict:
    url = SUMMARY if day is None else f"{SUMMARY}?target_date={day}"
    response = client.get(url)
    assert response.status_code == 200, response.text
    return response.json()


def test_resumen_consolida_por_vendedor_y_cuenta_estados(
    client, create_order, db, data, default_seller
):
    driver = data.user_vehicle.id_user_vehicle
    p0, p1, p2 = data.products
    a = create_order(items=[item(p0, 2), item(p1, 1)], id_user_vehicle=driver)  # 12000
    b = create_order(items=[item(p0, 1, "2500.00")], id_user_vehicle=driver)  # 2500
    c = create_order(items=[item(p2, 1)])  # 8000, sin conductor
    create_order(id_user_vehicle=driver)  # queda pendiente
    cancelled = create_order(id_user_vehicle=driver)
    _dispatch(client, a)
    _dispatch(client, b, deliver=True)
    _dispatch(client, c)
    db.get(Order, cancelled["id_order"]).state = OrderState.CANCELLED
    db.commit()

    body = _summary(client)

    assert body["closure_date"] == today_bogota().isoformat()
    cards = body["cards"]
    assert (cards["total_orders"], cards["pending"], cards["cancelled"]) == (5, 1, 1)
    assert (cards["dispatched"], cards["delivered"]) == (2, 1)
    assert cards["sellers_count"] == 2
    assert Decimal(cards["total_income"]) == Decimal(22500)
    assert body["exported"] is False

    default_row, driver_row = body["sellers"]  # ordenados por código: 9, 99
    assert (default_row["seller_code"], default_row["seller_name"]) == (9, "Hielo Foca")
    assert default_row["orders_count"] == default_row["unassigned_orders"] == 1
    assert default_row["products"] == [
        {"product_code": "Z93", "product_name": "BOTELLON AGUA", "quantity": 1}
    ]
    assert Decimal(default_row["total_sales"]) == Decimal(8000)

    assert (driver_row["seller_code"], driver_row["seller_name"]) == (
        99,
        "Carlos Conductor",
    )
    assert (driver_row["orders_count"], driver_row["unassigned_orders"]) == (2, 0)
    # Z91 a dos precios distintos se suma en una sola cantidad
    assert [(p["product_code"], p["quantity"]) for p in driver_row["products"]] == [
        ("Z91", 3),
        ("Z92", 1),
    ]
    assert driver_row["total_quantity"] == 4
    assert Decimal(driver_row["total_sales"]) == Decimal(14500)
    assert driver_row["exported_at"] is None


def test_sin_conductor_se_suma_a_los_pedidos_propios_del_vendedor_por_defecto(
    client, create_order, data, default_seller
):
    own = create_order(id_user_vehicle=default_seller.id_user_vehicle)  # 2 x Z91
    unassigned = create_order()  # 2 x Z91, sin conductor
    _dispatch(client, own)
    _dispatch(client, unassigned)

    (row,) = _summary(client)["sellers"]

    assert row["id_user_vehicle"] == default_seller.id_user_vehicle
    assert (row["orders_count"], row["unassigned_orders"]) == (2, 1)
    assert row["products"][0]["quantity"] == 4
    assert Decimal(row["total_sales"]) == Decimal(12000)


def test_sin_vendedor_por_defecto_responde_422_solo_si_hace_falta(
    client, create_order, data, monkeypatch
):
    monkeypatch.setattr(settings, "DEFAULT_SELLER_CODE", "98")  # no existe
    _dispatch(client, create_order(id_user_vehicle=data.user_vehicle.id_user_vehicle))

    assert client.get(SUMMARY).status_code == 200  # todos tienen conductor

    _dispatch(client, create_order())
    response = client.get(SUMMARY)

    assert response.status_code == 422
    assert response.json()["code"] == "DEFAULT_SELLER_MISSING"


def test_resumen_de_otro_dia_y_fecha_futura(client, create_order, db, data):
    yesterday = today_bogota() - timedelta(days=1)
    order = create_order(id_user_vehicle=data.user_vehicle.id_user_vehicle)
    _dispatch(client, order)
    _move(db, order, yesterday)

    assert _summary(client)["sellers"] == []
    assert _summary(client, yesterday)["cards"]["dispatched"] == 1

    response = client.get(f"{SUMMARY}?target_date={today_bogota() + timedelta(days=1)}")
    assert response.status_code == 422
    assert response.json()["code"] == "INVALID_CLOSURE_DATE"


def test_dia_sin_pedidos(client):
    body = _summary(client, date(2020, 1, 15))

    assert body["sellers"] == []
    assert body["exported"] is False
    assert body["cards"]["total_orders"] == 0
    assert Decimal(body["cards"]["total_income"]) == 0


def test_estado_de_exportacion_y_cambios_posteriores(client, create_order, db, data):
    driver = data.user_vehicle.id_user_vehicle
    _dispatch(client, create_order(id_user_vehicle=driver))  # 6000
    db.add(
        ShiftClosure(
            closure_date=today_bogota(),
            id_user_vehicle=driver,
            total_orders=1,
            total_sales=Decimal(6000),
            exported_siigo=True,
            exported_at=datetime.now(UTC),
        )
    )
    db.commit()

    body = _summary(client)
    assert body["exported"] is True
    assert body["sellers"][0]["exported_at"] is not None
    assert body["sellers"][0]["export_outdated"] is False

    _dispatch(client, create_order(id_user_vehicle=driver))  # llega otro pedido

    body = _summary(client)
    assert body["exported"] is False
    assert body["sellers"][0]["export_outdated"] is True


def test_codigo_de_vendedor_no_numerico_responde_422(client, create_order, db, data):
    data.user_vehicle.code = "AB"  # dato antiguo: el schema ya no lo permite
    db.commit()
    _dispatch(client, create_order(id_user_vehicle=data.user_vehicle.id_user_vehicle))

    response = client.get(SUMMARY)

    assert response.status_code == 422
    assert response.json()["code"] == "INVALID_SELLER_CODE"


def test_el_resumen_es_del_gestor_y_del_admin(admin):
    assert admin.get(SUMMARY).status_code == 200


# --- Exportación a Siigo (C3): POST /shifts_closures/{fecha}/export -----------------


def _export(client, day: date | None = None):
    return client.post(f"{URL}/{day or today_bogota()}/export")


def _sheet_rows(response) -> list[tuple]:
    assert response.status_code == 200, response.text
    sheet = load_workbook(BytesIO(response.content)).active
    return list(sheet.iter_rows(values_only=True))  # type: ignore


def _closures(db, day: date | None = None) -> dict[int, ShiftClosure]:
    return shift_closure_repository.get_by_date(db, day or today_bogota())


def test_exporta_excel_con_un_renglon_por_vendedor_producto_y_precio(
    client, create_order, data, default_seller
):
    driver = data.user_vehicle.id_user_vehicle
    p0, p1, p2 = data.products
    first = create_order(items=[item(p0, 2), item(p1, 1)], id_user_vehicle=driver)
    second = create_order(items=[item(p0, 1, "2500.00")], id_user_vehicle=driver)
    unassigned = create_order(items=[item(p2, 1)])
    _dispatch(client, first)
    _dispatch(client, second)
    _dispatch(client, unassigned, deliver=True)
    create_order(id_user_vehicle=driver)  # pendiente: no se factura

    response = _export(client)

    assert response.headers["content-type"] == MEDIA_TYPE
    assert f"cierre_{today_bogota()}.xlsx" in response.headers["content-disposition"]
    header, *rows = _sheet_rows(response)
    assert header[:3] == ("Fecha", "Cód. vendedor", "Vendedor")
    assert [r[1:] for r in rows] == [
        # vendedor por defecto (09 -> 9): el pedido sin conductor
        (9, "Hielo Foca", "Z93", "BOTELLON AGUA", 1, 8000, 8000),
        (None, "Hielo Foca", None, "Total vendedor", 1, None, 8000),
        # mismo producto a dos precios: renglones separados
        (99, "Carlos Conductor", "Z91", "BOLSA HIELO 2KG", 1, 2500, 2500),
        (99, "Carlos Conductor", "Z91", "BOLSA HIELO 2KG", 2, 3000, 6000),
        (99, "Carlos Conductor", "Z92", "BOLSA HIELO 5KG", 1, 6000, 6000),
        (None, "Carlos Conductor", None, "Total vendedor", 4, None, 14500),
        (None, None, None, "TOTAL DEL DÍA", 5, None, 22500),
    ]
    assert rows[0][0].date() == today_bogota()


def test_exportar_registra_los_cierres_y_el_resumen_queda_exportado(
    client, create_order, db, data, default_seller
):
    driver = data.user_vehicle.id_user_vehicle
    _dispatch(client, create_order(id_user_vehicle=driver))  # 6000
    _dispatch(client, create_order())  # 6000, sin conductor -> vendedor por defecto

    assert _export(client).status_code == 200

    closures = _closures(db)
    assert set(closures) == {driver, default_seller.id_user_vehicle}
    assert all(c.exported_siigo and c.exported_at for c in closures.values())
    assert (closures[driver].total_orders, closures[driver].total_sales) == (
        1,
        Decimal(6000),
    )
    body = _summary(client)
    assert body["exported"] is True
    assert all(row["exported_at"] for row in body["sellers"])


def test_reexportar_actualiza_el_mismo_cierre(client, create_order, db, data):
    driver = data.user_vehicle.id_user_vehicle
    _dispatch(client, create_order(id_user_vehicle=driver))
    _export(client)
    _dispatch(client, create_order(id_user_vehicle=driver))  # llega otro pedido
    assert _summary(client)["sellers"][0]["export_outdated"] is True

    assert _export(client).status_code == 200

    closures = _closures(db)
    assert len(closures) == 1
    assert (closures[driver].total_orders, closures[driver].total_sales) == (
        2,
        Decimal(12000),
    )
    assert _summary(client)["exported"] is True


def test_reexportar_quita_el_cierre_de_un_vendedor_sin_pedidos(
    client, create_order, db, data, default_seller
):
    driver = data.user_vehicle.id_user_vehicle
    order = create_order(id_user_vehicle=driver)
    _dispatch(client, order)
    _export(client)
    # se corrige el conductor después de exportar
    db.get(Order, order["id_order"]).id_user_vehicle = default_seller.id_user_vehicle
    db.commit()
    assert _summary(client)["exported"] is False

    _export(client)

    assert set(_closures(db)) == {default_seller.id_user_vehicle}
    assert _summary(client)["exported"] is True


def test_exportar_un_dia_anterior(client, create_order, db, data):
    yesterday = today_bogota() - timedelta(days=1)
    order = create_order(id_user_vehicle=data.user_vehicle.id_user_vehicle)
    _dispatch(client, order)
    _move(db, order, yesterday)

    rows = _sheet_rows(_export(client, yesterday))

    assert rows[1][0].date() == yesterday
    assert set(_closures(db, yesterday)) == {data.user_vehicle.id_user_vehicle}
    assert _closures(db) == {}


def test_exportar_sin_pedidos_facturables_responde_422(client, create_order, db):
    create_order()  # pendiente

    response = _export(client)

    assert response.status_code == 422
    assert response.json()["code"] == "NO_BILLABLE_ORDERS"
    assert _closures(db) == {}


def test_exportar_fecha_futura_responde_422(client):
    response = _export(client, today_bogota() + timedelta(days=1))

    assert response.status_code == 422
    assert response.json()["code"] == "INVALID_CLOSURE_DATE"


def test_exportar_sin_vendedor_por_defecto_no_registra_nada(
    client, create_order, db, monkeypatch
):
    monkeypatch.setattr(settings, "DEFAULT_SELLER_CODE", "98")  # no existe
    _dispatch(client, create_order())

    response = _export(client)

    assert response.status_code == 422
    assert response.json()["code"] == "DEFAULT_SELLER_MISSING"
    assert _closures(db) == {}
